import html
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .redaction import markdown_inline, redact_text

logger = logging.getLogger(__name__)


class Reporter:
    def __init__(self, report_dir: str = "reports/"):
        self.report_dir = Path(report_dir)
        self.report_dir.mkdir(parents=True, exist_ok=True)

    def print_summary(self, results: list[dict], stats: dict) -> None:

        try:
            from rich import box
            from rich.console import Console
            from rich.panel import Panel
            from rich.table import Table

            console = Console()

            ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
            console.print(
                Panel(
                    f"[bold red]LLM JAILBREAK FRAMEWORK[/bold red]\n"
                    f"[dim]Experiment Report — {ts}[/dim]",
                    border_style="red",
                )
            )

            console.print("\n[bold]OVERVIEW[/bold]")
            console.print(f"  Total Attacks:    [cyan]{stats.get('total_attacks', 0)}[/cyan]")
            console.print(f"  Successful:       [red]{stats.get('total_success', 0)}[/red]")
            sr = stats.get("success_rate", 0)
            color = "red" if sr > 0.5 else "yellow" if sr > 0.2 else "green"
            console.print(f"  Success Rate:     [{color}]{sr:.1%}[/{color}]")
            console.print(
                f"  Avg Risk Score:   [magenta]{stats.get('average_risk_score', 0):.3f}[/magenta]"
            )

            metadata = results[0].get("run_metadata") if results else None
            if metadata:
                console.print("\n[bold]RUN METADATA[/bold]")
                console.print(f"  Provider:       [cyan]{metadata.get('provider', '-')}[/cyan]")
                console.print(f"  Model:          [cyan]{metadata.get('model_name', '-')}[/cyan]")
                console.print(f"  Eval Mode:      [cyan]{metadata.get('eval_mode', '-')}[/cyan]")
                console.print(f"  Seed:           [cyan]{metadata.get('seed', '-')}[/cyan]")

            console.print("\n[bold]BY STRATEGY[/bold]")
            table = Table(box=box.ROUNDED, show_header=True, header_style="bold magenta")
            table.add_column("Strategy", style="cyan")
            table.add_column("Total", justify="right")
            table.add_column("Success", justify="right")
            table.add_column("Success Rate", justify="right")
            table.add_column("Avg Risk Score", justify="right")

            for strat, data in stats.get("by_strategy", {}).items():
                sr_val = data.get("success_rate", 0)
                color = "red" if sr_val > 0.5 else "yellow" if sr_val > 0.2 else "green"
                table.add_row(
                    strat,
                    str(data["total"]),
                    str(data["success"]),
                    f"[{color}]{sr_val:.1%}[/{color}]",
                    f"{data.get('avg_risk_score', 0):.3f}",
                )

            console.print(table)

            console.print("\n[bold]BY RISK LEVEL[/bold]")
            risk_colors = {
                "Critical": "bold red",
                "High": "red",
                "Medium": "yellow",
                "Low": "blue",
                "None": "green",
            }
            for level, count in stats.get("by_risk_level", {}).items():
                if count > 0:
                    color = risk_colors.get(level, "white")
                    console.print(f"  {level:10s}: [{color}]{count}[/{color}]")

            console.print("")

        except ImportError:
            self._plain_summary(stats)

    def _plain_summary(self, stats: dict) -> None:

        print("\n" + "=" * 60)
        print("JAILBREAK FRAMEWORK — EXPERIMENT RESULTS")
        print("=" * 60)
        print(f"Total Attacks:  {stats.get('total_attacks', 0)}")
        print(f"Successful:     {stats.get('total_success', 0)}")
        print(f"Success Rate:   {stats.get('success_rate', 0):.1%}")
        print(f"Avg Risk:       {stats.get('average_risk_score', 0):.3f}")
        print("=" * 60)

    def save_json_summary(self, stats: dict, session_id: str = "") -> Path:

        try:
            if not isinstance(stats, dict):
                logger.warning(f"Invalid stats type: {type(stats)}. Using empty dict.")
                stats = {}

            ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            path = self.report_dir / f"summary_{ts}.json"

            summary = {
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "session_id": session_id,
                "stats": stats,
            }

            with open(path, "w", encoding="utf-8") as f:
                json.dump(summary, f, indent=2)

            logger.info(f"JSON summary saved: {path}")
            return path
        except Exception as e:
            logger.error(f"Failed to save JSON summary: {e}")
            raise

    def save_markdown_report(self, results: list[dict], stats: dict, session_id: str = "") -> Path:

        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        path = self.report_dir / f"report_{ts}.md"

        lines = [
            "# LLM Jailbreak Framework — Experiment Report",
            "",
            f"**Generated:** {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}  ",
            f"**Session:** `{session_id}`",
            "",
            "---",
            "",
            "## Overview",
            "",
            "| Metric | Value |",
            "|--------|-------|",
            f"| Total Attacks | {stats.get('total_attacks', 0)} |",
            f"| Successful Jailbreaks | {stats.get('total_success', 0)} |",
            f"| Success Rate | {stats.get('success_rate', 0):.1%} |",
            f"| Average Risk Score | {stats.get('average_risk_score', 0):.4f} |",
            "",
            "## Results by Strategy",
            "",
            "| Strategy | Total | Success | Rate | Avg Risk |",
            "|----------|-------|---------|------|----------|",
        ]

        for strat, data in stats.get("by_strategy", {}).items():
            lines.append(
                f"| {strat} | {data['total']} | {data['success']} | "
                f"{data.get('success_rate', 0):.1%} | "
                f"{data.get('avg_risk_score', 0):.3f} |"
            )

        lines += [
            "",
            "## Risk Level Distribution",
            "",
            "| Risk Level | Count |",
            "|------------|-------|",
        ]
        for level, count in stats.get("by_risk_level", {}).items():
            lines.append(f"| {level} | {count} |")

        lines += [
            "",
            "## Sample Results",
            "",
        ]

        top = sorted(results, key=lambda r: r.get("risk_score", 0), reverse=True)[:5]
        for i, r in enumerate(top, 1):
            lines += [
                f"### Attack #{i} — {r.get('strategy', 'unknown')} [{r.get('risk_level', '?')}]",
                f"**Risk Score:** `{r.get('risk_score', 0):.4f}`  ",
                f"**Success:** `{r.get('success', False)}`  ",
                f"**Prompt (truncated):** `{markdown_inline(r.get('adversarial_prompt', '')[:150])}...`  ",
                f"**Response (truncated):** `{markdown_inline(r.get('raw_response', '')[:200])}...`  ",
                f"**Reasoning:** {markdown_inline(r.get('eval_reasoning', ''))}",
                "",
            ]

        lines += [
            "---",
            "",
            "*Generated by LLM Jailbreak Automation Framework — For research use only.*",
        ]

        path.write_text("\n".join(lines), encoding="utf-8")
        logger.info(f"Markdown report saved: {path}")
        return path

    def generate_charts(self, results: list[dict], stats: dict) -> list[Path]:

        try:
            import matplotlib.pyplot as plt
        except ImportError:
            logger.warning("matplotlib not installed. Skipping chart generation.")
            return []

        paths = []
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

        fig, axes = plt.subplots(1, 2, figsize=(14, 5))
        fig.suptitle("LLM Jailbreak Experiment Results", fontsize=14, fontweight="bold")
        fig.patch.set_facecolor("#1a1a2e")

        for ax in axes:
            ax.set_facecolor("#16213e")
            ax.tick_params(colors="white")
            ax.title.set_color("white")
            for spine in ax.spines.values():
                spine.set_color("#444")

        strat_data = stats.get("by_strategy", {})
        if strat_data:
            strategies = list(strat_data.keys())
            rates = [strat_data[s].get("success_rate", 0) for s in strategies]
            colors = ["#e63946" if r > 0.5 else "#f4a261" if r > 0.2 else "#2a9d8f" for r in rates]
            axes[0].bar(strategies, rates, color=colors, edgecolor="#333")
            axes[0].set_title("Success Rate by Strategy", color="white")
            axes[0].set_ylabel("Success Rate", color="white")
            axes[0].set_ylim(0, 1)
            for bar, val in zip(axes[0].patches, rates, strict=True):
                axes[0].text(
                    bar.get_x() + bar.get_width() / 2,
                    bar.get_height() + 0.02,
                    f"{val:.0%}",
                    ha="center",
                    color="white",
                    fontsize=9,
                )

        risk_data = stats.get("by_risk_level", {})
        if risk_data:
            level_colors = {
                "Critical": "#e63946",
                "High": "#f4a261",
                "Medium": "#e9c46a",
                "Low": "#2a9d8f",
                "None": "#457b9d",
            }
            labels = [k for k, v in risk_data.items() if v > 0]
            sizes = [risk_data[k] for k in labels]
            pie_colors = [level_colors.get(label, "#888") for label in labels]

            if sizes:
                wedges, texts, autotexts = axes[1].pie(
                    sizes,
                    labels=labels,
                    colors=pie_colors,
                    autopct="%1.1f%%",
                    startangle=90,
                    textprops={"color": "white"},
                )
                for at in autotexts:
                    at.set_color("white")
                axes[1].set_title("Risk Level Distribution", color="white")

        plt.tight_layout()
        chart_path = self.report_dir / f"chart_{ts}.png"
        plt.savefig(chart_path, dpi=150, bbox_inches="tight", facecolor="#1a1a2e", edgecolor="none")
        plt.close()
        paths.append(chart_path)
        logger.info(f"Charts saved: {chart_path}")

        return paths

    def save_html_dashboard(
        self,
        results: list[dict],
        stats: dict,
        session_id: str = "",
        comparison: dict[str, dict] | None = None,
        leaderboard: list[dict] | None = None,
    ) -> Path:
        """Self-contained interactive HTML dashboard (Plotly via CDN).

        `comparison` maps a model label to its own aggregate_stats() output,
        for a side-by-side multi-model view. `leaderboard` is the output of
        Scorer.leaderboard() — the "Top Failing Prompts" table.
        """

        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        path = self.report_dir / f"dashboard_{ts}.html"

        metadata = results[0].get("run_metadata") if results else None
        data: dict[str, Any] = {
            "stats": stats,
            "comparison": comparison or {},
            "leaderboard": [
                {
                    **row,
                    "base_prompt_text": redact_text(
                        row.get("base_prompt_text", ""), include_pii=True
                    )[:160],
                }
                for row in (leaderboard or [])
            ],
        }
        data_json = json.dumps(data, ensure_ascii=False)

        title = "Jailbreak Framework — Comparison Dashboard" if comparison else (
            "Jailbreak Framework — Experiment Dashboard"
        )
        generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        meta_html = ""
        if metadata:
            meta_html = (
                f"<p class='meta'>Provider: <b>{html.escape(str(metadata.get('provider', '-')))}</b>"
                f" &middot; Model: <b>{html.escape(str(metadata.get('model_name', '-')))}</b>"
                f" &middot; Eval mode: <b>{html.escape(str(metadata.get('eval_mode', '-')))}</b></p>"
            )

        page = f"""<!doctype html>
<html lang="en" data-theme="dark">
<head>
<meta charset="utf-8">
<title>{html.escape(title)}</title>
<script src="https://cdn.plot.ly/plotly-2.32.0.min.js"></script>
<style>
  :root {{ color-scheme: dark; }}
  body {{
    background: #12121f; color: #e6e6f0; font-family: -apple-system, Segoe UI, sans-serif;
    margin: 0; padding: 24px 32px 48px;
  }}
  h1 {{ color: #ff5c72; margin-bottom: 4px; }}
  .meta {{ color: #9a9ab0; margin-top: 0; }}
  .session {{ color: #6d6d85; font-size: 0.85em; }}
  .kpi-row {{ display: flex; gap: 16px; flex-wrap: wrap; margin: 20px 0 28px; }}
  .kpi {{
    background: #1b1b2e; border: 1px solid #2c2c44; border-radius: 10px;
    padding: 14px 20px; min-width: 150px;
  }}
  .kpi .label {{ font-size: 0.8em; color: #9a9ab0; text-transform: uppercase; letter-spacing: 0.04em; }}
  .kpi .value {{ font-size: 1.6em; font-weight: 600; margin-top: 4px; }}
  .charts {{ display: flex; flex-wrap: wrap; gap: 24px; }}
  .chart-box {{ flex: 1 1 460px; background: #1b1b2e; border-radius: 10px; padding: 8px; }}
  table {{ border-collapse: collapse; width: 100%; margin-top: 12px; }}
  th, td {{
    text-align: left; padding: 8px 10px; border-bottom: 1px solid #2c2c44; font-size: 0.9em;
  }}
  th {{ color: #9a9ab0; text-transform: uppercase; font-size: 0.75em; letter-spacing: 0.03em; }}
  section {{ margin-top: 36px; }}
  h2 {{ color: #e6e6f0; border-bottom: 1px solid #2c2c44; padding-bottom: 8px; }}
</style>
</head>
<body>
  <h1>{html.escape(title)}</h1>
  {meta_html}
  <p class="session">Generated {html.escape(generated)} &middot; Session {html.escape(str(session_id))}</p>

  <div class="kpi-row" id="kpi-row"></div>

  <div class="charts">
    <div class="chart-box" id="chart-strategy"></div>
    <div class="chart-box" id="chart-risk"></div>
  </div>

  <section id="comparison-section" style="display:none">
    <h2>Model Comparison</h2>
    <div class="chart-box" id="chart-comparison"></div>
    <table id="comparison-table"></table>
  </section>

  <section id="leaderboard-section" style="display:none">
    <h2>Top Failing Prompts</h2>
    <table id="leaderboard-table"></table>
  </section>

<script>
const DATA = {data_json};
const dark = {{
  paper_bgcolor: "#1b1b2e", plot_bgcolor: "#1b1b2e",
  font: {{ color: "#e6e6f0" }},
}};

function kpi(label, value) {{
  const div = document.createElement("div");
  div.className = "kpi";
  div.innerHTML = `<div class="label">${{label}}</div><div class="value">${{value}}</div>`;
  return div;
}}

const stats = DATA.stats || {{}};
const kpiRow = document.getElementById("kpi-row");
kpiRow.appendChild(kpi("Total Attacks", stats.total_attacks ?? 0));
kpiRow.appendChild(kpi("Successful", stats.total_success ?? 0));
kpiRow.appendChild(kpi("Success Rate", ((stats.success_rate ?? 0) * 100).toFixed(1) + "%"));
kpiRow.appendChild(kpi("Avg Risk Score", (stats.average_risk_score ?? 0).toFixed(3)));

const byStrategy = stats.by_strategy || {{}};
const strategies = Object.keys(byStrategy);
Plotly.newPlot("chart-strategy", [{{
  x: strategies,
  y: strategies.map(s => byStrategy[s].success_rate),
  type: "bar",
  marker: {{ color: strategies.map(s => byStrategy[s].success_rate > 0.5 ? "#e63946" : byStrategy[s].success_rate > 0.2 ? "#f4a261" : "#2a9d8f") }},
}}], {{ ...dark, title: "Success Rate by Strategy", yaxis: {{ tickformat: ".0%", range: [0, 1] }} }}, {{ responsive: true }});

const byRisk = stats.by_risk_level || {{}};
const riskLabels = Object.keys(byRisk).filter(k => byRisk[k] > 0);
const riskColors = {{ Critical: "#e63946", High: "#f4a261", Medium: "#e9c46a", Low: "#2a9d8f", None: "#457b9d" }};
Plotly.newPlot("chart-risk", [{{
  labels: riskLabels,
  values: riskLabels.map(k => byRisk[k]),
  type: "pie",
  marker: {{ colors: riskLabels.map(k => riskColors[k] || "#888") }},
}}], {{ ...dark, title: "Risk Level Distribution" }}, {{ responsive: true }});

const comparison = DATA.comparison || {{}};
const models = Object.keys(comparison);
if (models.length > 0) {{
  document.getElementById("comparison-section").style.display = "block";
  Plotly.newPlot("chart-comparison", [
    {{
      x: models, y: models.map(m => comparison[m].success_rate ?? 0),
      name: "Success Rate", type: "bar", marker: {{ color: "#e63946" }},
    }},
    {{
      x: models, y: models.map(m => comparison[m].average_risk_score ?? 0),
      name: "Avg Risk Score", type: "bar", marker: {{ color: "#f4a261" }}, yaxis: "y2",
    }},
  ], {{
    ...dark, title: "Model Comparison", barmode: "group",
    yaxis: {{ title: "Success Rate", tickformat: ".0%", range: [0, 1] }},
    yaxis2: {{ title: "Avg Risk Score", overlaying: "y", side: "right", range: [0, 1] }},
  }}, {{ responsive: true }});

  const compTable = document.getElementById("comparison-table");
  let rows = "<tr><th>Model</th><th>Total</th><th>Success</th><th>Success Rate</th><th>Avg Risk Score</th></tr>";
  for (const m of models) {{
    const s = comparison[m];
    rows += `<tr><td>${{m}}</td><td>${{s.total_attacks ?? 0}}</td><td>${{s.total_success ?? 0}}</td>` +
      `<td>${{((s.success_rate ?? 0) * 100).toFixed(1)}}%</td><td>${{(s.average_risk_score ?? 0).toFixed(3)}}</td></tr>`;
  }}
  compTable.innerHTML = rows;
}}

const leaderboard = DATA.leaderboard || [];
if (leaderboard.length > 0) {{
  document.getElementById("leaderboard-section").style.display = "block";
  const lbTable = document.getElementById("leaderboard-table");
  let rows = "<tr><th>#</th><th>Strategy</th><th>Category</th><th>Prompt</th><th>Success Rate</th><th>Avg Risk Score</th></tr>";
  leaderboard.forEach((row, i) => {{
    rows += `<tr><td>${{i + 1}}</td><td>${{row.strategy}}</td><td>${{row.category}}</td>` +
      `<td>${{row.base_prompt_text}}</td><td>${{((row.success_rate ?? 0) * 100).toFixed(1)}}%</td>` +
      `<td>${{(row.avg_risk_score ?? 0).toFixed(3)}}</td></tr>`;
  }});
  lbTable.innerHTML = rows;
}}
</script>
</body>
</html>
"""
        path.write_text(page, encoding="utf-8")
        logger.info(f"HTML dashboard saved: {path}")
        return path

    def save_pptx_summary(
        self,
        results: list[dict],
        stats: dict,
        session_id: str = "",
        chart_paths: list[Path] | None = None,
        comparison: dict[str, dict] | None = None,
    ) -> Path | None:
        """Portfolio-style slide deck: title, overview, chart image, per-strategy
        table, and (if provided) a per-model comparison table. Returns None
        (and logs a warning) if python-pptx isn't installed."""

        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError:
            logger.warning("python-pptx not installed. Skipping PPTX export.")
            return None

        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        path = self.report_dir / f"summary_{ts}.pptx"

        prs = Presentation()
        blank = prs.slide_layouts[6]

        # Title slide
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "LLM Jailbreak Framework"
        slide.placeholders[1].text = (
            f"Experiment Report — {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}\n"
            f"Session: {session_id}"
        )

        # Overview slide
        slide = prs.slides.add_slide(blank)
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
        tb.text_frame.text = "Overview"
        tb.text_frame.paragraphs[0].font.size = Pt(28)
        rows = [
            ("Total Attacks", str(stats.get("total_attacks", 0))),
            ("Successful Jailbreaks", str(stats.get("total_success", 0))),
            ("Success Rate", f"{stats.get('success_rate', 0):.1%}"),
            ("Average Risk Score", f"{stats.get('average_risk_score', 0):.4f}"),
        ]
        table_shape = slide.shapes.add_table(
            len(rows) + 1, 2, Inches(0.5), Inches(1.2), Inches(6), Inches(0.4 * (len(rows) + 1))
        )
        table = table_shape.table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        for i, (label, value) in enumerate(rows, start=1):
            table.cell(i, 0).text = label
            table.cell(i, 1).text = value

        # By-strategy slide
        by_strategy = stats.get("by_strategy", {})
        if by_strategy:
            slide = prs.slides.add_slide(blank)
            tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
            tb.text_frame.text = "Results by Strategy"
            tb.text_frame.paragraphs[0].font.size = Pt(28)
            strat_rows = list(by_strategy.items())
            table_shape = slide.shapes.add_table(
                len(strat_rows) + 1,
                4,
                Inches(0.5),
                Inches(1.2),
                Inches(8),
                Inches(0.35 * (len(strat_rows) + 1)),
            )
            table = table_shape.table
            for col, header in enumerate(["Strategy", "Total", "Success", "Success Rate"]):
                table.cell(0, col).text = header
            for i, (strat, data) in enumerate(strat_rows, start=1):
                table.cell(i, 0).text = strat
                table.cell(i, 1).text = str(data.get("total", 0))
                table.cell(i, 2).text = str(data.get("success", 0))
                table.cell(i, 3).text = f"{data.get('success_rate', 0):.1%}"

        # Comparison slide
        if comparison:
            slide = prs.slides.add_slide(blank)
            tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
            tb.text_frame.text = "Model Comparison"
            tb.text_frame.paragraphs[0].font.size = Pt(28)
            comp_rows = list(comparison.items())
            table_shape = slide.shapes.add_table(
                len(comp_rows) + 1,
                4,
                Inches(0.5),
                Inches(1.2),
                Inches(8),
                Inches(0.35 * (len(comp_rows) + 1)),
            )
            table = table_shape.table
            for col, header in enumerate(["Model", "Total", "Success Rate", "Avg Risk Score"]):
                table.cell(0, col).text = header
            for i, (model_label, model_stats) in enumerate(comp_rows, start=1):
                table.cell(i, 0).text = model_label
                table.cell(i, 1).text = str(model_stats.get("total_attacks", 0))
                table.cell(i, 2).text = f"{model_stats.get('success_rate', 0):.1%}"
                table.cell(i, 3).text = f"{model_stats.get('average_risk_score', 0):.4f}"

        # Chart image slide(s)
        for chart_path in chart_paths or []:
            if not Path(chart_path).exists():
                continue
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(chart_path), Inches(0.5), Inches(0.4), width=Inches(9)
            )

        prs.save(str(path))
        logger.info(f"PPTX summary saved: {path}")
        return path

    def save_pdf_summary(
        self,
        results: list[dict],
        stats: dict,
        session_id: str = "",
        chart_paths: list[Path] | None = None,
        comparison: dict[str, dict] | None = None,
    ) -> Path | None:
        """Single-file PDF report: overview, by-strategy and (optionally)
        per-model comparison tables, and embedded chart images. Returns None
        (and logs a warning) if reportlab isn't installed."""

        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import LETTER
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import (
                Image,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError:
            logger.warning("reportlab not installed. Skipping PDF export.")
            return None

        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        path = self.report_dir / f"summary_{ts}.pdf"

        styles = getSampleStyleSheet()
        story: list[Any] = []

        story.append(Paragraph("LLM Jailbreak Framework", styles["Title"]))
        story.append(
            Paragraph(
                f"Experiment Report — "
                f"{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
                styles["Normal"],
            )
        )
        story.append(Paragraph(f"Session: {html.escape(str(session_id))}", styles["Normal"]))
        story.append(Spacer(1, 0.3 * inch))

        def make_table(rows: list[list[str]]) -> Table:
            table = Table(rows, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1b1b2e")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.whitesmoke]),
                    ]
                )
            )
            return table

        story.append(Paragraph("Overview", styles["Heading2"]))
        story.append(
            make_table(
                [
                    ["Metric", "Value"],
                    ["Total Attacks", str(stats.get("total_attacks", 0))],
                    ["Successful Jailbreaks", str(stats.get("total_success", 0))],
                    ["Success Rate", f"{stats.get('success_rate', 0):.1%}"],
                    ["Average Risk Score", f"{stats.get('average_risk_score', 0):.4f}"],
                ]
            )
        )
        story.append(Spacer(1, 0.25 * inch))

        by_strategy = stats.get("by_strategy", {})
        if by_strategy:
            story.append(Paragraph("Results by Strategy", styles["Heading2"]))
            rows = [["Strategy", "Total", "Success", "Success Rate"]]
            for strat, data in by_strategy.items():
                rows.append(
                    [
                        strat,
                        str(data.get("total", 0)),
                        str(data.get("success", 0)),
                        f"{data.get('success_rate', 0):.1%}",
                    ]
                )
            story.append(make_table(rows))
            story.append(Spacer(1, 0.25 * inch))

        if comparison:
            story.append(Paragraph("Model Comparison", styles["Heading2"]))
            rows = [["Model", "Total", "Success Rate", "Avg Risk Score"]]
            for model_label, model_stats in comparison.items():
                rows.append(
                    [
                        model_label,
                        str(model_stats.get("total_attacks", 0)),
                        f"{model_stats.get('success_rate', 0):.1%}",
                        f"{model_stats.get('average_risk_score', 0):.4f}",
                    ]
                )
            story.append(make_table(rows))
            story.append(Spacer(1, 0.25 * inch))

        for chart_path in chart_paths or []:
            if not Path(chart_path).exists():
                continue
            story.append(Paragraph("Charts", styles["Heading2"]))
            story.append(Image(str(chart_path), width=6.5 * inch, height=2.7 * inch))
            story.append(Spacer(1, 0.2 * inch))

        doc = SimpleDocTemplate(str(path), pagesize=LETTER)
        doc.build(story)
        logger.info(f"PDF summary saved: {path}")
        return path
